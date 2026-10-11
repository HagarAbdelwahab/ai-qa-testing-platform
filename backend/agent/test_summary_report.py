"""
Test Summary Report agent — gathers cycle data from Jira + AIO, drafts scope from
epic/story descriptions with Claude, and renders the report as .docx and .pdf.

Two phases (the user edits everything in between):
  gather(cycle_name, keys, pull_aio)  -> prefill dict
  render(data, out_dir)               -> (docx_path, pdf_path)

CLI (used by the ~/.claude/skills/test-summary-report skill):
  python -m agent.test_summary_report <input.json> <output_dir>
"""
import json
import os
import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import arabic_reshaper
from bidi.algorithm import get_display

ARABIC = re.compile(r"[\u0600-\u06FF]")
FONT_CANDIDATES = [
    "/Library/Fonts/Arial Unicode.ttf",
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
]
HEADER_BG = "1F3864"
VERDICT_COLORS = {"GO": "2E7D32", "CONDITIONAL GO": "EF6C00", "NO-GO": "C62828"}
REPORTS_DIR = Path(__file__).parent.parent.parent / "reports" / "test-summary"


# ---------- data ----------

def normalise(d: dict) -> dict:
    """Coerce counts to ints and compute executed / execution rate / pass rate."""
    d = json.loads(json.dumps(d))
    d.setdefault("cycle", {})
    if not d.get("title"):
        d["title"] = "Test Summary Report" + (f" — {d['feature']}" if d.get("feature") else "")
    ex = d.setdefault("execution", {})
    for k in ("passed", "failed", "blocked", "not_run", "in_progress"):
        ex[k] = int(ex.get(k) or 0)
    total = int(ex.get("total") or sum(ex[k] for k in ("passed", "failed", "blocked", "not_run", "in_progress")))
    executed = ex["passed"] + ex["failed"] + ex["blocked"]
    ex["total"] = total
    ex["executed"] = int(ex.get("executed") or executed)
    ex["execution_rate"] = round(ex["executed"] / total * 100, 1) if total else 0
    ex["pass_rate"] = round(ex["passed"] / ex["executed"] * 100, 1) if ex["executed"] else 0
    return d


def sections(d: dict) -> list:
    """Normalised section list shared by both renderers: (heading, kind, payload)."""
    ex = d["execution"]
    out = []
    out.append(("Report Information", "kv", [
        ("Feature / Release", d.get("feature", "")),
        ("Test Cycle", f"{d['cycle'].get('name', '')}" + (f" (AIO ID {d['cycle']['id']})" if d["cycle"].get("id") else "")),
        ("Epic", f"{d.get('epic', {}).get('key', '')} — {d.get('epic', {}).get('summary', '')}".strip(" —")),
        ("Testing Period", f"{d.get('period', {}).get('start', '')} → {d.get('period', {}).get('end', '')}".strip(" →")),
        ("Environment", d.get("environment", "")),
        ("QA Owner(s)", ", ".join(d.get("qa_owners", []))),
        ("Prepared By", d.get("prepared_by", "")),
        ("Report Date", d.get("date", "")),
    ]))
    if d.get("objective"):
        out.append(("Objective", "text", d["objective"]))
    scope = d.get("scope", {})
    out.append(("Scope of Testing", "subsections", [
        ("In Scope", scope.get("in", [])),
        ("Out of Scope", scope.get("out", [])),
    ]))
    if d.get("stories"):
        out.append(("User Stories Covered", "table",
                    (["Key", "Summary", "Status"],
                     [[s.get("key", ""), s.get("summary", ""), s.get("status", "")] for s in d["stories"]])))
    src = "AIO cycle" if ex.get("source", "").lower() == "aio" else "manual input"
    out.append((f"Test Execution Summary (source: {src})", "table",
                (["Metric", "Count"], [
                    ["Total test cases in cycle", ex["total"]],
                    ["Executed", ex["executed"]],
                    ["Passed", ex["passed"]],
                    ["Failed", ex["failed"]],
                    ["Blocked", ex["blocked"]],
                    ["In Progress", ex["in_progress"]],
                    ["Not Run", ex["not_run"]],
                    ["Execution rate", f"{ex['execution_rate']}%"],
                    ["Pass rate (of executed)", f"{ex['pass_rate']}%"],
                ])))
    issues = d.get("known_issues", [])
    out.append(("Known Issues / Open Defects", "table" if issues else "text",
                (["Key", "Summary", "Severity", "Status"],
                 [[i.get("key", ""), i.get("summary", ""), i.get("severity", ""), i.get("status", "")] for i in issues])
                if issues else "No known open issues."))
    risks = d.get("risks", [])
    out.append(("Risks", "table" if risks else "text",
                (["Risk", "Impact", "Mitigation"],
                 [[r.get("risk", ""), r.get("impact", ""), r.get("mitigation", "")] for r in risks])
                if risks else "No risks identified."))
    concl = d.get("conclusion", {})
    out.append(("Conclusion & Recommendation", "verdict",
                (concl.get("recommendation", ""), concl.get("notes", ""))))
    return [(f"{n}. {h}", k, p) for n, (h, k, p) in enumerate(out, 1)]


# ---------- docx ----------

def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _docx_table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(str(h))
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(c, HEADER_BG)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
            if ARABIC.search(str(v)):
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def build_docx(d: dict, path: str):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    h = doc.add_heading(d.get("title", "Test Summary Report"), level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, kind, payload in sections(d):
        doc.add_heading(heading, level=1)
        if kind == "kv":
            _docx_table(doc, ["Field", "Value"], [[k, v] for k, v in payload if v])
        elif kind == "text":
            doc.add_paragraph(payload)
        elif kind == "subsections":
            for sub, items in payload:
                doc.add_heading(sub, level=2)
                if not items:
                    doc.add_paragraph("—")
                for it in items:
                    doc.add_paragraph(it, style="List Bullet")
        elif kind == "table":
            _docx_table(doc, *payload)
        elif kind == "verdict":
            verdict, notes = payload
            p = doc.add_paragraph()
            r = p.add_run(f"Recommendation: {verdict}")
            r.bold = True
            r.font.size = Pt(13)
            col = VERDICT_COLORS.get(verdict.upper(), "000000")
            r.font.color.rgb = RGBColor.from_string(col)
            if notes:
                doc.add_paragraph(notes)
    doc.save(path)


# ---------- pdf ----------

def _font():
    for f in FONT_CANDIDATES:
        if os.path.exists(f):
            pdfmetrics.registerFont(TTFont("Uni", f))
            return "Uni"
    return "Helvetica"


def _t(s) -> str:
    """Shape Arabic for reportlab (it has no bidi support) and escape markup."""
    s = str(s)
    if ARABIC.search(s):
        s = get_display(arabic_reshaper.reshape(s))
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(d: dict, path: str):
    font = _font()
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName=font, fontSize=9.5, leading=13)
    cellw = ParagraphStyle("cw", parent=body, textColor=colors.white)
    h1 = ParagraphStyle("h1", parent=ss["Heading2"], fontName=font, textColor=colors.HexColor("#" + HEADER_BG))
    h2 = ParagraphStyle("h2", parent=ss["Heading3"], fontName=font)
    title = ParagraphStyle("t", parent=ss["Title"], fontName=font, alignment=TA_CENTER)
    width = A4[0] - 4 * cm

    def table(header, rows):
        data = [[Paragraph(_t(h), cellw) for h in header]] + \
               [[Paragraph(_t(v), body) for v in r] for r in rows]
        t = Table(data, colWidths=[width / len(header)] * len(header) if len(header) != 2
                  else [width * 0.35, width * 0.65], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + HEADER_BG)),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F4F8")]),
        ]))
        return t

    flow = [Paragraph(_t(d.get("title", "Test Summary Report")), title), Spacer(1, 8)]
    for heading, kind, payload in sections(d):
        flow.append(Paragraph(_t(heading), h1))
        if kind == "kv":
            flow.append(table(["Field", "Value"], [[k, v] for k, v in payload if v]))
        elif kind == "text":
            flow.append(Paragraph(_t(payload), body))
        elif kind == "subsections":
            for sub, items in payload:
                flow.append(Paragraph(_t(sub), h2))
                for it in (items or ["—"]):
                    flow.append(Paragraph("• " + _t(it), body))
        elif kind == "table":
            flow.append(table(*payload))
        elif kind == "verdict":
            verdict, notes = payload
            col = VERDICT_COLORS.get(verdict.upper(), "000000")
            flow.append(Paragraph(f'<font color="#{col}" size="13"><b>Recommendation: {_t(verdict)}</b></font>', body))
            flow.append(Spacer(1, 6))
            if notes:
                flow.append(Paragraph(_t(notes), body))
        flow.append(Spacer(1, 6))
    SimpleDocTemplate(path, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                      title=d.get("title", "Test Summary Report")).build(flow)


def render(d: dict, out_dir: str | Path = REPORTS_DIR) -> tuple[str, str]:
    d = normalise(d)
    os.makedirs(out_dir, exist_ok=True)
    slug = re.sub(r"[^A-Za-z0-9]+", "_", d["cycle"].get("name", "") or "cycle").strip("_") or "cycle"
    base = os.path.join(str(out_dir), f"Test_Summary_Report_{slug}_{date.today():%Y%m%d}")
    build_docx(d, base + ".docx")
    build_pdf(d, base + ".pdf")
    return base + ".docx", base + ".pdf"


# ---------- gather ----------

def _child_items(epic_key: str) -> list[dict]:
    from jira_client import get_issues_by_jql
    return get_issues_by_jql(f"parent = {epic_key} AND issuetype != Bug ORDER BY key")


def _open_bugs(epic_keys: list[str], story_keys: list[str]) -> list[dict]:
    """Open bugs under the epic(s), sub-task bugs of the stories, and bugs linked to the stories."""
    from jira_client import get_issues_by_jql
    clauses = []
    if epic_keys:
        clauses.append(f"parent in ({', '.join(epic_keys)})")
    if story_keys:
        clauses.append(f"parent in ({', '.join(story_keys)})")
        clauses += [f'issue in linkedIssues("{k}")' for k in story_keys]
    if not clauses:
        return []
    found: dict[str, dict] = {}
    # Query clause by clause so one unsupported clause doesn't hide the rest.
    for c in clauses:
        jql = f"issuetype = Bug AND statusCategory != Done AND ({c}) ORDER BY priority DESC"
        try:
            for i in get_issues_by_jql(jql):
                f = i["fields"]
                found[i["key"]] = {
                    "key": i["key"],
                    "summary": f.get("summary", ""),
                    "severity": (f.get("priority") or {}).get("name", ""),
                    "status": (f.get("status") or {}).get("name", ""),
                }
        except Exception as e:
            print(f"  [jira] bug query failed for ({c}): {str(e).splitlines()[0]}")
    return list(found.values())


def _draft_scope(epic: dict, stories: list[dict]) -> dict:
    """Ask Claude for objective + in/out scope, grounded only in the descriptions."""
    from test_generator import _call_ai
    blocks = []
    if epic:
        blocks.append(f"EPIC {epic['key']}: {epic['summary']}\n{epic.get('description', '')[:6000]}")
    for s in stories:
        blocks.append(f"STORY {s['key']}: {s['summary']}\n{s.get('description', '')[:3000]}")
    prompt = f"""You are a senior QA lead writing the scope section of a Test Summary Report.

Use ONLY the epic/story text below. Do not invent features, rules or exclusions.

{chr(10).join(blocks)}

Return ONLY a JSON object with:
- "objective": 1-2 sentences describing what this testing cycle validated (from the epic goal, or the stories if no epic).
- "in_scope": 4-10 short bullets naming the functional areas, flows, validations, roles and integrations described above.
- "out_of_scope": bullets ONLY for items the text explicitly excludes or defers. Empty list if none are stated.

No explanation, JSON only."""
    raw = _call_ai(prompt, max_tokens=2000).strip()
    m = re.search(r"\{.*\}", raw, re.S)
    if not m:
        raise ValueError(f"Scope draft was not JSON: {raw[:200]}")
    d = json.loads(m.group())
    return {"objective": d.get("objective", ""), "in": d.get("in_scope", []), "out": d.get("out_of_scope", [])}


def gather(cycle_name: str, keys: list[str], pull_aio: bool = True) -> dict:
    """Collect everything the report can derive; the user fills/edits the rest."""
    from jira_client import get_issue, adf_to_text
    epic, stories, epic_keys = {}, [], []
    for k in keys:
        print(f"  [jira] fetching {k}...")
        f = get_issue(k)["fields"]
        item = {"key": k, "summary": f["summary"], "status": f["status"]["name"],
                "description": adf_to_text(f.get("description")) if f.get("description") else ""}
        if f["issuetype"]["name"] == "Epic":
            epic_keys.append(k)
            epic = epic or item
            children = _child_items(k)
            print(f"  [jira] {k} is an epic with {len(children)} child item(s)")
            for c in children:
                cf = c["fields"]
                ck = c["key"]
                if any(s["key"] == ck for s in stories):
                    continue
                cfull = get_issue(ck)["fields"]
                stories.append({"key": ck, "summary": cf["summary"], "status": cf["status"]["name"],
                                "description": adf_to_text(cfull.get("description")) if cfull.get("description") else ""})
        elif not any(s["key"] == k for s in stories):
            stories.append(item)

    story_keys = [s["key"] for s in stories]
    print(f"  [jira] {len(stories)} story(ies) in scope; looking up open bugs...")
    bugs = _open_bugs(epic_keys, story_keys)
    print(f"  [jira] {len(bugs)} open bug(s) found")

    print("  [ai] drafting scope from epic/story descriptions...")
    try:
        scope = _draft_scope(epic, stories)
    except Exception as e:
        print(f"  [ai] scope draft failed: {e}")
        scope = {"objective": "", "in": [], "out": []}

    execution = {"source": "manual"}
    aio_error = ""
    cycle = {"name": cycle_name}
    if pull_aio and cycle_name:
        print(f"  [aio] pulling execution stats for cycle '{cycle_name}'...")
        import aio_client, config
        st = aio_client.get_cycle_stats(config.JIRA_PROJECT_ID, cycle_name)
        if st.get("error"):
            aio_error = st["error"]
            print(f"  [aio] {aio_error}")
        else:
            cycle = {"name": st.get("cycle_title") or cycle_name, "id": st.get("cycle_id")}
            execution = {"source": "aio", **{k: st.get(k, 0) for k in
                         ("total", "passed", "failed", "blocked", "not_run", "in_progress")}}
            print(f"  [aio] {cycle['name']}: {st.get('total', 0)} case(s), {st.get('passed', 0)} passed")

    return {
        "title": "",
        "feature": epic.get("summary", ""),
        "cycle": cycle,
        "epic": {"key": epic["key"], "summary": epic["summary"]} if epic else {},
        "stories": [{k: s[k] for k in ("key", "summary", "status")} for s in stories],
        "objective": scope["objective"],
        "scope": {"in": scope["in"], "out": scope["out"]},
        "qa_owners": [],
        "period": {"start": "", "end": ""},
        "environment": "",
        "execution": execution,
        "aio_error": aio_error,
        "known_issues": bugs,
        "risks": [],
        "conclusion": {"recommendation": "", "notes": ""},
        "prepared_by": "",
        "date": f"{date.today():%Y-%m-%d}",
    }


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as f:
        d = json.load(f)
    for p in render(d, sys.argv[2]):
        print(p)


if __name__ == "__main__":
    main()
